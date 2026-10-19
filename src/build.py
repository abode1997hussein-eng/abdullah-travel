"""Build the site + the Excel price tabs from data.py.

python3 -I build.py [SHEET_ID]
  -> ../index.html          standalone site for Netlify / GitHub Pages
  -> ../artifact.html       same page for the Claude artifact preview
  -> ../أسعار-الموقع.xlsx    the two clean tabs to upload to Google Drive
"""
import json, os, re, sys
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import data  # noqa: E402

SHEET_ID = sys.argv[1] if len(sys.argv) > 1 else ""
SITE_URL = os.environ.get("SITE_URL", "https://abode1997hussein-eng.github.io/abdullah-travel/")
TODAY = date.today().strftime("%d-%m-%Y")

photos = sorted(f[:-5] for f in os.listdir(os.path.join(SITE, "assets/photos")) if f.endswith(".webp"))
flags = sorted(f[:-4] for f in os.listdir(os.path.join(SITE, "assets/flags")) if f.endswith(".svg"))
credits = json.load(open(os.path.join(SITE, "assets/photos/credits.json"), encoding="utf-8"))

used = sorted({g["الرمز"].lower() for g in data.GROUPS if g["الرمز"]} & set(photos))
credit_line = ""
if used:
    parts = [f'<a href="{credits[k]["page"]}" target="_blank" rel="noopener">{credits[k]["title"]}</a> '
             f'({credits[k]["artist"] or "مجهول"}، {credits[k]["license"]})' for k in used if k in credits]
    credit_line = '<p>صور المعالم من ويكيميديا كومنز: ' + "، ".join(parts) + '</p>'

payload = {"offers": data.GROUPS, "visas": data.VISAS}
tpl = open(os.path.join(HERE, "template.html"), encoding="utf-8").read()
page = (tpl.replace("{{DATA}}", json.dumps(payload, ensure_ascii=False))
           .replace("{{PHOTOS}}", json.dumps(photos)).replace("{{FLAGS}}", json.dumps(flags))
           .replace("{{PHOTO_DIR}}", "assets/photos/").replace("{{FLAG_DIR}}", "assets/flags/")
           .replace("{{HERO}}", "img/hero.jpg").replace("{{AVATAR}}", "img/avatar.jpg")
           .replace("{{UPDATED}}", TODAY).replace("{{CREDITS}}", credit_line)
           .replace("{{SHEET_ID}}", SHEET_ID))
left = re.findall(r"\{\{\w+\}\}", page)
assert not left, left

open(os.path.join(SITE, "artifact.html"), "w", encoding="utf-8").write(page)

# ---- standalone document
title = re.search(r"<title>(.*?)</title>", page).group(1)
body = re.sub(r"<title>.*?</title>\s*", "", page, count=1)
body = re.sub(r'<meta name="claude:viewport-fit-sides"[^>]*>\s*', "", body, count=1)
font = re.search(r'<link rel="stylesheet" href="https://fonts\.googleapis\.com[^>]*>', body).group(0)
body = body.replace(font + "\n", "", 1)
style_end = body.index("</style>") + len("</style>")
style, rest = body[:style_end], body[style_end:]
DESC = ("عبدالله حسين علوان للسفر والسياحة، بغداد: كروبات سياحية، حفلات ورأس السنة، "
        "فيزا لأكثر من 35 دولة، وتذاكر طيران. الحجز مباشرة على الواتساب.")
doc = f'''<!doctype html>
<html lang="ar" dir="rtl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">
<title>{title} — سفر وسياحة</title>
<meta name="description" content="{DESC}">
<meta name="theme-color" content="#100E0C">
<meta property="og:type" content="website">
<meta property="og:locale" content="ar_IQ">
<meta property="og:title" content="{title} — سفر وسياحة">
<meta property="og:description" content="{DESC}">
<meta property="og:url" content="{SITE_URL}">
<meta property="og:image" content="{SITE_URL}img/og.jpg">
<meta property="og:image:width" content="1200">
<meta property="og:image:height" content="630">
<meta name="twitter:card" content="summary_large_image">
<meta name="twitter:image" content="{SITE_URL}img/og.jpg">
<link rel="canonical" href="{SITE_URL}">
<link rel="icon" href="img/avatar.jpg">
<link rel="apple-touch-icon" href="img/avatar.jpg">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
{font}
<style>
:root{{color-scheme:light;padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}}
body{{margin:0}} img{{max-width:100%}} [hidden]{{display:none!important}}
</style>
{style}
</head>
<body>{rest.strip()}
</body>
</html>
'''
open(os.path.join(SITE, "index.html"), "w", encoding="utf-8").write(doc)

# ---- Excel tabs
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AR = "Arial"
thin = Side(style="thin", color="D9D2CA")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()
ws = wb.active
ws.title = "تعليمات"
ws.sheet_view.rightToLeft = True
lines = [
    ("الموقع يقرا التبويبين «كروبات» و«فيزا» من هذا الملف مباشرة", True),
    ("أي تعديل تسويه يظهر بالموقع خلال دقائق، بدون ما أحد يعدّل الموقع.", False),
    ("", False),
    ("القواعد", True),
    ("١ — لا تغيّر اسم التبويبين، ولا صف العناوين الأول، ولا ترتيب الأعمدة.", False),
    ("٢ — «ظاهر»: نعم حتى يطلع بالموقع، لا حتى ينخفي بدون ما تحذف السطر.", False),
    ("٣ — العرض الواحد ممكن ياخذ أكثر من سطر (سطر لكل فندق). اكتب «الوجهة» بكل سطر بنفس الكتابة بالضبط، وباقي التفاصيل بالسطر الأول بس.", False),
    ("٤ — «النوع»: كروب أو حفلة. الحفلات تطلع بقسم الحفلات.", False),
    ("٥ — «المنطقة» هي الفلتر فوق الكروبات (تركيا، لبنان، مصر، الخليج، آسيا…). أي اسم جديد تكتبه يصير فلتر جديد.", False),
    ("٦ — «الرمز»: رمز المطار بالإنكليزي، ويطلع على الصورة. الصور المتوفرة حالياً: " + "، ".join(p.upper() for p in photos) + ".", False),
    ("٧ — الأسعار: 1075$ بالدولار، أو 520,000 د.ع بالدينار. «سعر البيع» فقط — لا تكتب الكلفة أو العمولة.", False),
    ("٨ — القوائم داخل خانة وحدة (التواريخ، يشمل، المتطلبات) افصل بينها بفاصلة منقوطة ؛", False),
    ("٩ — إذا كتبت أكثر من تاريخ مفصولة بـ ؛ يطلع للزبون يختار منها. تاريخ واحد يطلع نص عادي.", False),
    ("١٠ — بالفيزا: «العلم» رمز الدولة بحرفين (ae, th, tr…). الأعلام المتوفرة: " + "، ".join(flags) + ".", False),
    ("", False),
    ("الخانات الصفراء هي الأسعار — أكثر شي راح تعدّله.", False),
]
for i, (t, bold) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=AR, bold=bold, size=12 if bold else 10)
    c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
ws.column_dimensions["A"].width = 110

def sheet(name, head, rows, price_cols):
    s = wb.create_sheet(name)
    s.sheet_view.rightToLeft = True
    s.append(head)
    for r in rows:
        s.append([r.get(h, "") for h in head])
    for ci, h in enumerate(head, 1):
        c = s.cell(row=1, column=ci)
        c.font = Font(name=AR, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="100E0C")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = B
        longest = max(len(str(s.cell(row=r, column=ci).value or "")) for r in range(1, s.max_row + 1))
        s.column_dimensions[get_column_letter(ci)].width = min(max(longest + 2, 10), 42)
    for r in range(2, s.max_row + 1):
        for ci, h in enumerate(head, 1):
            c = s.cell(row=r, column=ci)
            c.font = Font(name=AR, size=10)
            c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            c.border = B
            if h in price_cols:
                c.fill = PatternFill("solid", fgColor="FDF1D6")
    s.freeze_panes = "E2" if name == "كروبات" else "D2"
    s.row_dimensions[1].height = 30
    s.auto_filter.ref = f"A1:{get_column_letter(len(head))}{s.max_row}"
    return s

sheet("كروبات", data.G_HEAD, data.GROUPS, {"سعر البيع", "السنكل", "الطفل بدون سرير", "الرضيع"})
sheet("فيزا", data.V_HEAD, data.VISAS, {"سعر البيع"})
xlsx = os.path.join(SITE, "أسعار-الموقع.xlsx")
wb.save(xlsx)

print("offers rows:", len(data.GROUPS), "| visa rows:", len(data.VISAS))
print("photos:", photos)
print("sheet id:", SHEET_ID or "(none — fallback data only)")
print("wrote index.html, artifact.html,", os.path.basename(xlsx))
